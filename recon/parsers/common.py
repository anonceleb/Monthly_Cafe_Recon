from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl


class ParseError(Exception):
    """Raised when a file does not have the fixed layout we expect."""


def num(v, default=Decimal("0")) -> Decimal:
    if v is None:
        return default
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "").replace("₹", "").replace("%", "")
    if s in ("", "-", "NA", "nan", "None"):
        return default
    try:
        return Decimal(s)
    except InvalidOperation:
        return default


def text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s if s and s.lower() not in ("nan", "none") else None


def norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def to_dt(v) -> datetime:
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%y %H:%M", "%d/%m/%Y %H:%M:%S",
                "%d/%m/%Y %I:%M:%S %p", "%Y-%m-%d", "%d/%b/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ParseError(f"unrecognised date/time {v!r}")


def to_date(v) -> date | None:
    return to_dt(v).date() if v not in (None, "", "-") else None


def sheet_rows(path: Path, sheet: str | None = None) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def sheet_names(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return wb.sheetnames


def header_map(row: list, required: dict[str, str], where: str) -> dict[str, int]:
    """Map logical names to column index by (normalised) header prefix; fail loudly if the layout moved."""
    normed = [norm(c) for c in row]
    out, missing = {}, []
    for key, label in required.items():
        want = norm(label)
        idx = next((i for i, h in enumerate(normed) if h.startswith(want)), None)
        if idx is None:
            missing.append(label)
        else:
            out[key] = idx
    if missing:
        raise ParseError(f"{where}: expected column(s) not found: {missing}. The report layout may have changed.")
    return out


def find_row(rows: list[list], predicate, limit: int = 60) -> int | None:
    for i, r in enumerate(rows[:limit]):
        if predicate(r):
            return i
    return None


def stated_period(rows: list[list], label: str, fmts=("%Y-%m-%d", "%d %b %Y", "%d/%m/%Y", "%d/%m/%y")):
    """Read a printed 'Report period' cell such as '01 Apr 2026 - 30 Apr 2026' or '2026-04-01 - 2026-06-30'. None if absent/unreadable."""
    for r in rows[:25]:
        cells = [str(c).strip() for c in r if c not in (None, "")]
        for i, c in enumerate(cells):
            if c.lower().startswith(label.lower()):
                blob = " ".join(cells[i:])
                found = re.findall(r"\d{4}-\d{2}-\d{2}|\d{1,2} [A-Za-z]{3} \d{4}|\d{2}/\d{2}/\d{2,4}", blob)
                if len(found) >= 2:
                    out = []
                    for f in found[:2]:
                        for fmt in fmts:
                            try:
                                out.append(datetime.strptime(f, fmt).date()); break
                            except ValueError:
                                pass
                    if len(out) == 2:
                        return out[0], out[1]
    return None


def snap_month(lo: date, hi: date, days: int = 3):
    """Files without a printed period are monthly exports: if the rows start/stop within `days` of a month edge, extend to that edge."""
    import calendar
    if lo.day <= days:
        lo = lo.replace(day=1)
    last = calendar.monthrange(hi.year, hi.month)[1]
    if last - hi.day < days:
        hi = hi.replace(day=last)
    return lo, hi

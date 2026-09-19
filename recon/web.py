"""Upload page + dashboard + findings, served locally. Every upload is ingested, then the rules re-run over everything held."""
from __future__ import annotations

import io
import json
import tempfile
from decimal import Decimal
from pathlib import Path

from fastapi import FastAPI, File, Request, UploadFile
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
import openpyxl

from . import config as cfgmod, db, engine, ingest
from .parsers import LABELS

app = FastAPI(title="Kredo reconciliation")
templates = Jinja2Templates(directory=str(Path(__file__).with_name("templates")))

SEVERITY_ORDER = {"gap": 0, "warn": 1, "info": 2}
SOURCES = [  # what a complete monthly close needs, in the order people gather it
    ("bank_icici", "ICICI bank statement", "Detailed statement, xlsx"),
    ("pos_bills", "Petpooja order summary", "Order_Summary_Item_Report csv"),
    ("pos_online", "Petpooja online-platform sales", "Sales Report: Online Platforms xlsx"),
    ("zomato_delivery", "Zomato delivery settlement", "Settlement Report xlsx"),
    ("zomato_dining", "Zomato dining payout", "Consolidated payout report xlsx"),
    ("swiggy_orders", "Swiggy orders annexure", "consolidate-annexure-orders csv"),
    ("swiggy_adjustments", "Swiggy adjustments", "consolidate-annexure-adjustment csv"),
]


def inr(v, dp=2):
    """Indian digit grouping: 1,23,456.78"""
    if v is None or v == "":
        return "—"
    v = Decimal(str(v))
    neg, v = v < 0, abs(v)
    whole, _, frac = f"{v:.{dp}f}".partition(".")
    head, tail = whole[:-3], whole[-3:]
    if head:
        head = ",".join([head[max(i - 2, 0):i] for i in range(len(head), 0, -2)][::-1])
        whole = head + "," + tail
    return ("-" if neg else "") + "₹" + whole + (f".{frac}" if dp else "")


templates.env.filters["inr"] = inr
templates.env.filters["inr0"] = lambda v: inr(v, 0)


def _latest_run(conn):
    return conn.execute("select * from run order by id desc limit 1").fetchone()


def _counts(conn, run_id):
    rows = conn.execute("select rule_id, severity, count(*) n, coalesce(sum(abs(diff)),0) amt from finding where run_id=%s group by 1,2", (run_id,)).fetchall()
    by_sev = {"gap": 0, "warn": 0, "info": 0}
    by_rule: dict = {}
    for r in rows:
        by_sev[r["severity"]] += r["n"]
        d = by_rule.setdefault(r["rule_id"], {"gap": 0, "warn": 0, "info": 0, "amt": Decimal(0)})
        d[r["severity"]] = r["n"]
        d["amt"] += r["amt"] if r["severity"] != "info" else 0
    return by_sev, by_rule


@app.get("/")
def dashboard(request: Request):
    with db.connect() as conn:
        db.init_schema(conn)
        run = _latest_run(conn)
        uploads = {r["kind"]: r for r in conn.execute(
            "select kind, count(*) files, min(period_from) lo, max(period_to) hi, max(uploaded_at) at from upload group by kind").fetchall()}
        ctx = {"request": request, "run": run, "uploads": uploads, "sources": SOURCES, "by_sev": None, "by_rule": {}, "summary": {}}
        if run:
            ctx["by_sev"], ctx["by_rule"] = _counts(conn, run["id"])
            ctx["summary"] = run["summary"]
            ctx["top"] = conn.execute("select * from finding where run_id=%s and severity='gap' order by abs(coalesce(diff,0)) desc limit 8", (run["id"],)).fetchall()
    return templates.TemplateResponse(request, "dashboard.html", ctx)


@app.get("/upload")
def upload_form(request: Request):
    with db.connect() as conn:
        history = conn.execute("select * from upload order by id desc limit 40").fetchall()
    return templates.TemplateResponse(request, "upload.html", {"history": history, "results": None, "labels": LABELS})


@app.post("/upload")
async def upload(request: Request, files: list[UploadFile] = File(...)):
    cfg = cfgmod.load()
    results = []
    with db.connect() as conn, tempfile.TemporaryDirectory() as tmp:
        db.init_schema(conn)
        for f in files:
            if not f.filename:
                continue
            path = Path(tmp) / Path(f.filename).name
            path.write_bytes(await f.read())
            try:
                results.append(ingest.ingest(conn, path, f.filename, cfg))
            except Exception as e:                       # a bad file must not sink the batch
                results.append({"filename": f.filename, "kind": None, "label": None, "status": "rejected", "rows_read": 0, "rows_new": 0, "note": f"Could not read this file: {e}"})
        ran = any(r["status"] == "loaded" for r in results)
        run_id = engine.run(conn, cfg) if ran else None
        history = conn.execute("select * from upload order by id desc limit 40").fetchall()
    return templates.TemplateResponse(request, "upload.html", {"history": history, "results": results, "ran": ran, "run_id": run_id, "labels": LABELS})


@app.post("/run")
def rerun():
    with db.connect() as conn:
        engine.run(conn)
    return RedirectResponse("/", status_code=303)


@app.get("/findings")
def findings(request: Request, rule: str = "", severity: str = "", platform: str = ""):
    with db.connect() as conn:
        run = _latest_run(conn)
        rows, rules = [], {}
        if run:
            where, args = ["run_id=%s"], [run["id"]]
            for col, val in (("rule_id", rule), ("severity", severity), ("platform", platform)):
                if val:
                    where.append(f"{col}=%s"); args.append(val)
            rows = conn.execute(f"select * from finding where {' and '.join(where)} order by case severity when 'gap' then 0 when 'warn' then 1 else 2 end, rule_id, abs(coalesce(diff,0)) desc limit 600", args).fetchall()
            rules = run["summary"].get("rules", {})
    return templates.TemplateResponse(request, "findings.html", {"rows": rows, "rules": rules, "f": {"rule": rule, "severity": severity, "platform": platform}, "run": run})


@app.get("/rules")
def rules_page(request: Request):
    cfg = cfgmod.load()
    from .rules import RULES
    terms = {f"{p} {pr}": v for (p, pr), v in cfg.terms.items()}
    return templates.TemplateResponse(request, "rules.html", {"rules": sorted(RULES, key=lambda r: r.id), "cfg": cfg, "terms": terms})


@app.get("/export.xlsx")
def export():
    with db.connect() as conn:
        run = _latest_run(conn)
        if not run:
            return RedirectResponse("/", status_code=303)
        rows = conn.execute("select * from finding where run_id=%s order by case severity when 'gap' then 0 when 'warn' then 1 else 2 end, rule_id, abs(coalesce(diff,0)) desc", (run["id"],)).fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Findings"
    ws.append(["Severity", "Rule", "Platform", "Product", "Reference", "Date", "Expected", "Actual", "Difference", "Message"])
    for r in rows:
        ws.append([r["severity"], r["rule_id"], r["platform"], r["product"], r["ref"], r["on_date"], r["expected"], r["actual"], r["diff"], r["message"]])
    ws2 = wb.create_sheet("Platform flow")
    flow = run["summary"].get("flow", [])
    if flow:
        ws2.append(list(flow[0].keys()))
        for f in flow:
            ws2.append(list(f.values()))
    ws3 = wb.create_sheet("Settlements")
    st = run["summary"].get("settlement", {})
    ws3.append(["Platform:product", "Settlements", "Matched", "Via adjustment", "Period boundary", "Gap", "Report expects", "Bank received", "Gap amount"])
    for k, v in st.items():
        ws3.append([k, v["utrs"], v["matched"], v["via_adjustment"], v["boundary"], v["gap"], round(v["expected"], 2), round(v["bank"], 2), round(v["gap_amount"], 2)])
    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col[:60]) + 2, 90)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="reconciliation-run-{run["id"]}.xlsx"'})

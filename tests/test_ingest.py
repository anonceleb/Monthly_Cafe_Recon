"""Parsers and ingest against the real April 2026 files (skipped when they are not on this machine)."""
from pathlib import Path

import openpyxl
import pytest

from recon import ingest
from recon.parsers import detect

APRIL = Path("/Users/ashwin/Streak/L'Amour/Kredo /April 2026")
pytestmark = pytest.mark.skipif(not APRIL.exists(), reason="April 2026 files not present")

EXPECT = {
    "Bank statements/2026-03-16 to 2026-05-15 Cafe ICICI Statement.xlsx": ("bank_icici", 1494),
    "Petpooja/Order_Summary_Item_Report_152749_2026-04-01_2026-04-30.csv": ("pos_bills", 1268),
    "Petpooja/Sales_Report_Online_Platforms_2026_09_18_15_51_09.xlsx": ("pos_online", 230),
    "Zomato/Zomato_Settlement_Report_21273585_01 Apr 2026_30 Apr 2026.xlsx": ("zomato_delivery", 131),
    "Zomato/Consolidated_payout_report_2026-04-01_to_2026-06-30.xlsx": ("zomato_dining", 559 + 40),
    "Swiggy/consolidate-annexure-orders_39cd23bd-9b82-4939-beff-427536bb0be81.csv": ("swiggy_orders", 228),
    "Swiggy/consolidate-annexure-adjustment_39cd23bd-9b82-4939-beff-427536bb0be81.csv": ("swiggy_adjustments", 4),
}


@pytest.mark.parametrize("rel,expect", EXPECT.items())
def test_each_file_is_recognised_and_fully_loaded(conn, cfg, rel, expect):
    kind, rows = expect
    assert detect(APRIL / rel) == kind
    r = ingest.ingest(conn, APRIL / rel, cfg=cfg)
    assert (r["status"], r["rows_read"], r["rows_new"]) == ("loaded", rows, rows)     # nothing silently collapsed (bank Tran. Id and refund ids repeat)


def test_same_file_twice_is_a_noop(conn, cfg):
    f = APRIL / "Bank statements/2026-03-16 to 2026-05-15 Cafe ICICI Statement.xlsx"
    ingest.ingest(conn, f, cfg=cfg)
    assert ingest.ingest(conn, f, cfg=cfg)["status"] == "duplicate"
    assert conn.execute("select count(*) n from bank_txn").fetchone()["n"] == 1494


def test_overlapping_statement_adds_only_new_rows(conn, cfg, tmp_path):
    f = APRIL / "Bank statements/2026-03-16 to 2026-05-15 Cafe ICICI Statement.xlsx"
    wb = openpyxl.load_workbook(f)
    ws = wb.active
    last = next(i for i in range(ws.max_row, 0, -1) if str(ws.cell(i, 1).value).strip().isdigit())
    for i in range(last, 17 + 800, -1):                     # keep the first 800 transactions only
        ws.delete_rows(i)
    part = tmp_path / "partial.xlsx"; wb.save(part)
    a = ingest.ingest(conn, part, cfg=cfg)
    b = ingest.ingest(conn, f, cfg=cfg)                     # the full statement overlaps the first 800
    assert a["rows_new"] == 800 and b["rows_new"] == 1494 - 800
    assert conn.execute("select count(*) n from bank_txn").fetchone()["n"] == 1494


def test_unrecognised_and_ignored_files_say_why(conn, cfg, tmp_path):
    junk = tmp_path / "x.csv"; junk.write_text("a,b\n1,2\n")
    assert ingest.ingest(conn, junk, cfg=cfg)["status"] == "rejected"
    r = ingest.ingest(conn, APRIL / "Petpooja/Supplier_Report_2026_09_14_06_15_49.xlsx", cfg=cfg)
    assert r["status"] == "ignored" and "not part of payout" in r["note"]


def test_april_end_to_end_headline_numbers(conn, cfg):
    from recon import engine
    for rel in EXPECT:
        ingest.ingest(conn, APRIL / rel, cfg=cfg)
    rid = engine.run(conn, cfg)
    s = conn.execute("select summary from run where id=%s", (rid,)).fetchone()["summary"]["settlement"]
    assert s["zomato:dining"]["matched"] == 12 and s["swiggy:dining"]["matched"] == 8
    assert s["swiggy:delivery"]["matched"] == 5 and s["swiggy:delivery"]["via_adjustment"] == 4   # weekly ads explained
    assert s["zomato:delivery"]["boundary"] == 2 and s["zomato:delivery"]["gap"] == 0
    pi4 = conn.execute("select sum(diff) d from finding where run_id=%s and rule_id='PI-04'", (rid,)).fetchone()["d"]
    assert float(pi4) == pytest.approx(-26372.64, abs=0.5)
